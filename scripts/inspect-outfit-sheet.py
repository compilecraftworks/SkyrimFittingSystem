"""Read every worksheet and merged-cell boundary without editing the source."""
import argparse
import hashlib
import json
from pathlib import Path
import sys

import openpyxl

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('workbook', type=Path)
parser.add_argument('output', type=Path)
args = parser.parse_args()
sys.stdout.reconfigure(encoding='utf-8')
wb = openpyxl.load_workbook(args.workbook, read_only=False, data_only=False)
out = {'source': str(args.workbook), 'sha256': hashlib.sha256(args.workbook.read_bytes()).hexdigest(), 'sheets': []}
for sheet in wb:
    rows = []
    for row in sheet.iter_rows():
        cells = [{'cell': c.coordinate, 'value': c.value, 'type': c.data_type} for c in row if c.value is not None]
        if cells:
            rows.append(cells)
    out['sheets'].append({'name': sheet.title, 'maxRow': sheet.max_row, 'maxCol': sheet.max_column,
                          'merges': [str(m) for m in sheet.merged_cells.ranges], 'rows': rows})
    print(json.dumps({'sheet': sheet.title, 'size': [sheet.max_row, sheet.max_column], 'rows': len(rows)}, ensure_ascii=False))
args.output.write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
